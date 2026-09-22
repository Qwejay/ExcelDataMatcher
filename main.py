import os
import sys
import re
import csv
import datetime
import threading
import queue
import subprocess
import platform
import tkinter as tk
from tkinter import filedialog, messagebox
import ttkbootstrap as ttk
from ttkbootstrap.constants import *
from ttkbootstrap.tooltip import ToolTip
from openpyxl import load_workbook, Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
import xlrd
import traceback

if platform.system() == "Windows":
    try:
        from ctypes import windll
        windll.shcore.SetProcessDpiAwareness(1)
    except Exception:
        pass

try:
    import windnd
    HAS_WINDND = True
except ImportError:
    HAS_WINDND = False

try:
    from tkinterdnd2 import DND_FILES
    HAS_TKDND = True
except ImportError:
    HAS_TKDND = False

__app_name__ = "ExcelDataMatcher"
__version__ = "3.0.0"
__author__ = "Qwejayhuang"
__copyright__ = f"Copyright © 2026 {__author__}"
__description__ = "Excel 表格数据智能提取工具"


def sanitize_excel_value(val):
    if val is None:
        return ""
    if isinstance(val, str):
        return ILLEGAL_CHARACTERS_RE.sub("", val)
    return val


class ExcelExtractorApp:
    def __init__(self, root):
        self.root = root
        self.root.title(f"{__app_name__} v{__version__} —— {__description__}")

        self.apply_screen_geometry()
        self.set_app_icon()

        self.file_path = None
        self.sheet_names = []
        self.extracted_headers = []
        self.extracted_data = []
        self.is_processing = False
        self.is_advanced_expanded = False

        self.stat_scanned_var = tk.StringVar(value="0")
        self.stat_matched_var = tk.StringVar(value="0")
        self.stat_ratio_var = tk.StringVar(value="0.0%")
        self.status_var = tk.StringVar(value="就绪：请先选择或拖入 Excel 文件")

        self.selected_sheet = tk.StringVar()
        self.search_all_sheets = tk.BooleanVar(value=True)
        self.header_row = tk.StringVar(value="1")
        self.no_header = tk.BooleanVar(value=False)
        
        self.sort_by_keyword_order = tk.BooleanVar(value=True)
        self.remove_duplicates_var = tk.BooleanVar(value=False)
        self.remove_inner_empty_var = tk.BooleanVar(value=True)
        self.remove_empty_rows_var = tk.BooleanVar(value=True)
        self.include_source_sheet_var = tk.BooleanVar(value=True)

        self.match_mode = tk.StringVar(value="contains")
        self.exclude_mode = tk.BooleanVar(value=False)
        self.case_sensitive = tk.BooleanVar(value=False)
        self.target_columns = tk.StringVar(value="")

        self.msg_queue = queue.Queue()

        self.create_modern_layout()
        self.create_context_menu()
        self.setup_drag_and_drop()

        self.toggle_search_all()
        self.toggle_no_header()
        
        self.root.protocol("WM_DELETE_WINDOW", self.on_closing)
        self.root.after(80, self.process_queue)

    def apply_screen_geometry(self):
        screen_w = self.root.winfo_screenwidth()
        screen_h = self.root.winfo_screenheight()
        win_w = min(1120, max(940, int(screen_w * 0.82)))
        win_h = min(730, max(580, int(screen_h * 0.82)))
        pos_x = max(0, int((screen_w - win_w) / 2))
        pos_y = max(0, int((screen_h - win_h) / 2) - 20)
        self.root.geometry(f"{win_w}x{win_h}+{pos_x}+{pos_y}")
        self.root.minsize(920, 560)

    def set_app_icon(self):
        icon_name = "logo.ico"
        if getattr(sys, 'frozen', False) and hasattr(sys, '_MEIPASS'):
            icon_path = os.path.join(sys._MEIPASS, icon_name)
        else:
            icon_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), icon_name)

        if os.path.exists(icon_path):
            try:
                self.root.iconbitmap(icon_path)
            except Exception:
                try:
                    icon_img = tk.PhotoImage(file=icon_path)
                    self.root.tk.call('wm', 'iconphoto', self.root._w, icon_img)
                except Exception:
                    pass

    def setup_drag_and_drop(self):
        if HAS_WINDND:
            try:
                windnd.hook_dropfiles(self.root, func=self.on_drop_files_safe)
            except Exception as e:
                print(f"windnd 挂载提示: {e}")
        elif HAS_TKDND and hasattr(self.root, 'drop_target_register'):
            try:
                self.root.drop_target_register(DND_FILES)
                self.root.dnd_bind('<<Drop>>', lambda e: self.on_drop_files_safe(self.root.tk.splitlist(e.data)))
            except Exception:
                pass

    def on_drop_files_safe(self, files):
        if not files: return
        try:
            file_path = files[0]
            if isinstance(file_path, bytes):
                file_path = file_path.decode('gbk', errors='ignore')
            self.msg_queue.put(("SAFE_DROP_EVENT", file_path))
        except Exception as e:
            print(f"拖拽数据传递异常: {e}")

    def create_modern_layout(self):
        main_box = ttk.Frame(self.root, padding=(10, 10, 10, 8))
        main_box.pack(fill=tk.BOTH, expand=True)

        sidebar = ttk.Frame(main_box, width=350)
        sidebar.pack(side=tk.LEFT, fill=tk.Y, padx=(0, 10))
        sidebar.pack_propagate(False)

        btn_area = ttk.Frame(sidebar)
        btn_area.pack(fill=tk.X, side=tk.BOTTOM, pady=(6, 0))

        self.btn_run = ttk.Button(btn_area, text=" 🚀 开始提取数据 ", command=self.start_extract_thread, bootstyle="success", padding=8)
        self.btn_run.pack(fill=tk.X, pady=(0, 5))

        self.save_btn = ttk.Button(btn_area, text=" 💾 导出数据 (Excel / CSV) ", command=self.save_to_file, bootstyle="info-outline", state=tk.DISABLED, padding=6)
        self.save_btn.pack(fill=tk.X)

        self.adv_toggle_btn = ttk.Button(
            sidebar, 
            text="⚙️ 高级选项 ▾ (点击展开)", 
            command=self.toggle_advanced_panel, 
            bootstyle="secondary-link"
        )
        self.adv_toggle_btn.pack(fill=tk.X, side=tk.BOTTOM, pady=(4, 2))

        self.adv_container = ttk.LabelFrame(sidebar, text=" 高级选项 ", padding=6)
        self.build_advanced_widgets()

        source_card = ttk.LabelFrame(sidebar, text=" 数据源 ", padding=(8, 6))
        source_card.pack(fill=tk.X, side=tk.TOP, pady=(0, 6))

        top_file_row = ttk.Frame(source_card)
        top_file_row.pack(fill=tk.X, pady=(0, 4))
        
        self.file_status_label = ttk.Label(top_file_row, text="点击或将文件拖入窗口", font=("Microsoft YaHei", 9, "bold"))
        self.file_status_label.pack(side=tk.LEFT, fill=tk.X, expand=True)

        ttk.Button(top_file_row, text="浏览...", command=self.select_file, bootstyle="primary-outline", width=8).pack(side=tk.RIGHT)

        sheet_row = ttk.Frame(source_card)
        sheet_row.pack(fill=tk.X)
        ttk.Label(sheet_row, text="Sheet:").pack(side=tk.LEFT)
        self.sheet_combobox = ttk.Combobox(sheet_row, textvariable=self.selected_sheet, state='readonly', width=13)
        self.sheet_combobox.pack(side=tk.LEFT, padx=6)
        ttk.Checkbutton(sheet_row, text="全表", variable=self.search_all_sheets, command=self.toggle_search_all).pack(side=tk.RIGHT)

        kw_card = ttk.LabelFrame(sidebar, text=" 匹配关键词 / 正则列表 (每行一个) ", padding=6)
        kw_card.pack(fill=tk.BOTH, side=tk.TOP, expand=True)

        kw_toolbar = ttk.Frame(kw_card)
        kw_toolbar.pack(fill=tk.X, pady=(0, 2))
        self.kw_count_label = ttk.Label(kw_toolbar, text="词数: 0", font=("Microsoft YaHei", 8), bootstyle="secondary")
        self.kw_count_label.pack(side=tk.LEFT)
        ttk.Button(kw_toolbar, text="清空", command=self.clear_keywords, bootstyle="secondary-link").pack(side=tk.RIGHT)
        ttk.Button(kw_toolbar, text="📂 导入文件", command=self.import_keywords_file, bootstyle="primary-link").pack(side=tk.RIGHT, padx=4)

        self.column_text = tk.Text(kw_card, height=4, font=("Microsoft YaHei", 9), relief="solid", borderwidth=1)
        self.column_text.pack(fill=tk.BOTH, expand=True)
        self.column_text.bind("<KeyRelease>", self.update_kw_count)

        content_area = ttk.Frame(main_box)
        content_area.pack(side=tk.RIGHT, fill=tk.BOTH, expand=True)

        stats_strip = ttk.Frame(content_area, padding=(6, 4))
        stats_strip.pack(fill=tk.X, pady=(0, 6))

        def add_kpi_metric(parent, label_text, val_var, text_color="primary"):
            box = ttk.Frame(parent)
            box.pack(side=tk.LEFT, padx=(0, 22))
            ttk.Label(box, text=label_text, font=("Microsoft YaHei", 8), bootstyle="secondary").pack(anchor=tk.W)
            ttk.Label(box, textvariable=val_var, font=("Microsoft YaHei", 14, "bold"), bootstyle=text_color).pack(anchor=tk.W)

        add_kpi_metric(stats_strip, "已扫描总数据", self.stat_scanned_var, "secondary")
        add_kpi_metric(stats_strip, "匹配命中行数", self.stat_matched_var, "success")
        add_kpi_metric(stats_strip, "提取命中率", self.stat_ratio_var, "info")

        status_box = ttk.Frame(stats_strip)
        status_box.pack(side=tk.RIGHT, fill=tk.Y, pady=2)
        ttk.Label(status_box, text="系统状态", font=("Microsoft YaHei", 8), bootstyle="secondary").pack(anchor=tk.E)
        self.status_label = ttk.Label(status_box, textvariable=self.status_var, font=("Microsoft YaHei", 9, "bold"), bootstyle="primary")
        self.status_label.pack(anchor=tk.E)

        self.progress_bar = ttk.Progressbar(content_area, mode='indeterminate', bootstyle=STRIPED)
        self.progress_bar.pack(fill=tk.X, pady=(0, 6))

        table_card = ttk.LabelFrame(content_area, text=" 数据预览 (Top 100) ", padding=5)
        table_card.pack(fill=tk.BOTH, expand=True)

        self.tree = ttk.Treeview(table_card, show="headings", selectmode="extended")
        tree_scroll_y = ttk.Scrollbar(table_card, orient=tk.VERTICAL, command=self.tree.yview)
        tree_scroll_x = ttk.Scrollbar(table_card, orient=tk.HORIZONTAL, command=self.tree.xview)
        self.tree.configure(yscrollcommand=tree_scroll_y.set, xscrollcommand=tree_scroll_x.set)

        self.tree.grid(row=0, column=0, sticky="nsew")
        tree_scroll_y.grid(row=0, column=1, sticky="ns")
        tree_scroll_x.grid(row=1, column=0, sticky="ew")
        table_card.rowconfigure(0, weight=1)
        table_card.columnconfigure(0, weight=1)

        self.tree.bind("<Double-1>", self.on_treeview_double_click)

        footer = ttk.Frame(content_area)
        footer.pack(fill=tk.X, pady=(5, 0))
        ttk.Label(footer, text="💡 提示：双击单元格复制内容", font=("Microsoft YaHei", 8), bootstyle="secondary").pack(side=tk.LEFT)
        ttk.Label(footer, text=f"{__copyright__} · 保留所有权利", font=("Microsoft YaHei", 8), bootstyle="secondary").pack(side=tk.RIGHT)

    def build_advanced_widgets(self):
        r1 = ttk.Frame(self.adv_container)
        r1.pack(fill=tk.X, pady=2)
        ttk.Label(r1, text="模式:").pack(side=tk.LEFT)
        ttk.Radiobutton(r1, text="包含", value="contains", variable=self.match_mode).pack(side=tk.LEFT, padx=2)
        ttk.Radiobutton(r1, text="精确", value="exact", variable=self.match_mode).pack(side=tk.LEFT, padx=2)
        ttk.Radiobutton(r1, text="正则", value="regex", variable=self.match_mode).pack(side=tk.LEFT, padx=2)

        r2 = ttk.Frame(self.adv_container)
        r2.pack(fill=tk.X, pady=2)
        ttk.Label(r2, text="指定检索列:").pack(side=tk.LEFT)
        self.col_entry = ttk.Entry(r2, textvariable=self.target_columns, width=12)
        self.col_entry.pack(side=tk.LEFT, padx=5)
        ToolTip(self.col_entry, text="如：A, C 或 1, 3 (留空全表检索)")
        ttk.Checkbutton(r2, text="大小写敏感", variable=self.case_sensitive).pack(side=tk.RIGHT)

        r3 = ttk.Frame(self.adv_container)
        r3.pack(fill=tk.X, pady=2)
        ttk.Label(r3, text="表头所在行:").pack(side=tk.LEFT)
        self.header_entry = ttk.Entry(r3, textvariable=self.header_row, width=6)
        self.header_entry.pack(side=tk.LEFT, padx=5)
        ttk.Checkbutton(r3, text="无表头", variable=self.no_header, command=self.toggle_no_header).pack(side=tk.LEFT)
        ttk.Checkbutton(r3, text="反向黑名单排除", variable=self.exclude_mode, bootstyle="danger").pack(side=tk.RIGHT)

        ttk.Separator(self.adv_container, orient=tk.HORIZONTAL).pack(fill=tk.X, pady=3)

        r4 = ttk.Frame(self.adv_container)
        r4.pack(fill=tk.X, pady=1)
        ttk.Checkbutton(r4, text="去单元格换行", variable=self.remove_inner_empty_var).pack(side=tk.LEFT, expand=True, anchor=tk.W)
        ttk.Checkbutton(r4, text="移除空行", variable=self.remove_empty_rows_var).pack(side=tk.LEFT, expand=True, anchor=tk.W)

        r5 = ttk.Frame(self.adv_container)
        r5.pack(fill=tk.X, pady=1)
        ttk.Checkbutton(r5, text="去重匹配行", variable=self.remove_duplicates_var).pack(side=tk.LEFT, expand=True, anchor=tk.W)
        ttk.Checkbutton(r5, text="保留来源Sheet", variable=self.include_source_sheet_var).pack(side=tk.LEFT, expand=True, anchor=tk.W)

        r6 = ttk.Frame(self.adv_container)
        r6.pack(fill=tk.X, pady=1)
        ttk.Checkbutton(r6, text="按关键词顺序排序 (默认推荐)", variable=self.sort_by_keyword_order, bootstyle="primary").pack(side=tk.LEFT, expand=True, anchor=tk.W)

    def toggle_advanced_panel(self):
        self.is_advanced_expanded = not self.is_advanced_expanded
        if self.is_advanced_expanded:
            self.adv_container.pack(fill=tk.X, side=tk.BOTTOM, before=self.adv_toggle_btn, pady=(0, 4))
            self.adv_toggle_btn.config(text="⚙️ 收起高级选项 ▴", bootstyle="secondary-outline")
        else:
            self.adv_container.pack_forget()
            self.adv_toggle_btn.config(text="⚙️ 高级选项 ▾ (点击展开)", bootstyle="secondary-link")

    def toggle_search_all(self):
        self.sheet_combobox.config(state='disabled' if self.search_all_sheets.get() else 'readonly')

    def toggle_no_header(self):
        self.header_entry.config(state=tk.DISABLED if self.no_header.get() else tk.NORMAL)

    def create_context_menu(self):
        self.context_menu = tk.Menu(self.root, tearoff=0)
        self.context_menu.add_command(label="粘贴", command=lambda: self.column_text.event_generate("<<Paste>>"))
        self.context_menu.add_command(label="复制", command=lambda: self.column_text.event_generate("<<Copy>>"))
        self.context_menu.add_command(label="清空", command=self.clear_keywords)
        self.column_text.bind("<Button-3>", lambda e: self.context_menu.post(e.x_root, e.y_root))

    def update_kw_count(self, event=None):
        lines = [line.strip() for line in self.column_text.get("1.0", "end-1c").splitlines() if line.strip()]
        self.kw_count_label.config(text=f"词数: {len(lines)}")

    def clear_keywords(self):
        self.column_text.delete("1.0", tk.END)
        self.update_kw_count()

    def import_keywords_file(self):
        file_path = filedialog.askopenfilename(
            title="选择关键词文件", 
            filetypes=[("文本或CSV文件", "*.txt *.csv"), ("所有文件", "*.*")]
        )
        if not file_path: return

        try:
            with open(file_path, "r", encoding="utf-8-sig", errors="ignore") as f:
                lines = [line.strip() for line in f if line.strip()]
            
            if lines:
                self.column_text.delete("1.0", tk.END)
                self.column_text.insert(tk.END, "\n".join(lines))
                self.update_kw_count()
                messagebox.showinfo("导入成功", f"成功载入 {len(lines)} 个关键词！")
            else:
                messagebox.showwarning("提示", "所选文件内容为空！")
        except Exception as e:
            messagebox.showerror("导入失败", f"无法读取关键词文件:\n{str(e)}")

    def on_treeview_double_click(self, event):
        region = self.tree.identify_region(event.x, event.y)
        if region != "cell": return
        column_id = self.tree.identify_column(event.x)
        item_id = self.tree.identify_row(event.y)
        if not item_id or not column_id: return

        col_index = int(column_id.replace("#", "")) - 1
        item_values = self.tree.item(item_id, "values")
        if item_values and col_index < len(item_values):
            cell_val = str(item_values[col_index])
            self.root.clipboard_clear()
            self.root.clipboard_append(cell_val)
            self.status_var.set(f"已复制: \"{cell_val[:20]}{'...' if len(cell_val) > 20 else ''}\"")
            self.status_label.configure(bootstyle="info")

    def select_file(self):
        if self.is_processing:
            messagebox.showwarning("警告", "后台正在提取数据，请勿切换文件！")
            return

        file_path = filedialog.askopenfilename(filetypes=[("Excel 文件", "*.xlsx *.xls")])
        if file_path:
            self.load_excel_file(file_path)

    def load_excel_file(self, file_path):
        wb = None
        try:
            if file_path.endswith(".xlsx"):
                wb = load_workbook(file_path, read_only=True)
                self.sheet_names = wb.sheetnames
            else:
                wb = xlrd.open_workbook(file_path)
                self.sheet_names = wb.sheet_names()
            
            if not self.sheet_names:
                raise ValueError("该 Excel 文件中没有包含任何有效的工作表！")

            self.file_path = file_path
            short_name = os.path.basename(file_path)
            display_title = short_name[:18] + ("..." if len(short_name) > 18 else "")
            self.file_status_label.config(text=f"✔ {display_title}", bootstyle="primary")
            
            self.sheet_combobox['values'] = self.sheet_names
            self.selected_sheet.set(self.sheet_names[0])
            self.toggle_search_all()
            
            self.status_var.set("文件读取成功，请配置规则")
            self.status_label.configure(bootstyle="success")
        except Exception as e:
            self.status_var.set("文件解析失败！")
            self.status_label.configure(bootstyle="danger")
            messagebox.showerror("文件读取失败", f"无法解析该文件:\n{str(e)}")
            self.file_path = None
            self.file_status_label.config(text="点击或将文件拖入窗口", bootstyle="default")
            self.sheet_combobox['values'] = []
            self.selected_sheet.set("")
        finally:
            if wb and hasattr(wb, 'close'):
                try: wb.close()
                except Exception: pass

    def parse_target_columns(self):
        raw = self.target_columns.get().strip()
        if not raw: return None
        cols = []
        for item in re.split(r'[,，\s]+', raw):
            if not item: continue
            if item.isdigit():
                val = int(item) - 1
                if val >= 0: cols.append(val)
            elif item.isalpha():
                idx = 0
                for char in item.upper():
                    idx = idx * 26 + (ord(char) - ord('A')) + 1
                cols.append(idx - 1)
        return list(set(cols)) if cols else None

    def parse_header_row(self):
        if self.no_header.get(): return None
        try:
            val = int(self.header_row.get().strip())
            return max(0, val - 1)
        except ValueError:
            return 0 

    def start_extract_thread(self):
        if self.is_processing: return

        if not self.file_path or not os.path.exists(self.file_path):
            messagebox.showwarning("提示", "请先选择或拖入有效的 Excel 文件！")
            return

        keywords = [col.strip() for col in self.column_text.get("1.0", "end-1c").splitlines() if col.strip()]
        if not keywords:
            messagebox.showwarning("提示", "请输入至少一个关键词或正则表达规则！")
            return

        match_mode = self.match_mode.get()
        if match_mode == "regex":
            for kw in keywords:
                try: re.compile(kw)
                except re.error as e:
                    messagebox.showerror("正则表达式语法错误", f"规则 [{kw}] 不合法:\n{str(e)}")
                    return

        config = {
            "file_path": self.file_path,
            "sheet_names": list(self.sheet_names),
            "search_all": self.search_all_sheets.get(),
            "selected_sheet": self.selected_sheet.get(),
            "header_idx": self.parse_header_row(),
            "target_cols": self.parse_target_columns(),
            "keywords": keywords,
            "match_mode": match_mode,
            "case_sensitive": self.case_sensitive.get(),
            "exclude_mode": self.exclude_mode.get(),
            "sort_by_kw": self.sort_by_keyword_order.get(),
            "remove_inner_empty": self.remove_inner_empty_var.get(),
            "remove_empty_rows": self.remove_empty_rows_var.get(),
            "remove_duplicates": self.remove_duplicates_var.get(),
            "include_source_sheet": self.include_source_sheet_var.get()
        }

        self.is_processing = True
        self.btn_run.config(state=tk.DISABLED)
        self.save_btn.config(state=tk.DISABLED)
        self.status_label.configure(bootstyle="primary")
        self.progress_bar.start(10)
        self.status_var.set("正在飞速检索分析中...")
        self.tree.delete(*self.tree.get_children()) 

        threading.Thread(target=self.run_fast_extraction, args=(config,), daemon=True).start()

    def run_fast_extraction(self, cfg):
        try:
            file_path = cfg["file_path"]
            sheets_to_read = cfg["sheet_names"] if cfg["search_all"] else [cfg["selected_sheet"]]
            header_idx = cfg["header_idx"]
            target_cols = cfg["target_cols"]
            keywords = cfg["keywords"]
            mode = cfg["match_mode"]
            case_sen = cfg["case_sensitive"]
            exclude_mode = cfg["exclude_mode"]

            regex_patterns = []
            if mode == "regex":
                flags = 0 if case_sen else re.IGNORECASE
                regex_patterns = [re.compile(kw, flags) for kw in keywords]
            elif not case_sen:
                keywords_processed = [kw.lower() for kw in keywords]
            else:
                keywords_processed = keywords

            sheets_data = self.read_all_sheets_data(file_path, sheets_to_read)

            extracted_items = []
            headers = []
            total_scanned_rows = 0
            order_counter = 0

            for sheet_name, rows in sheets_data.items():
                if not rows: continue

                if header_idx is not None and header_idx < len(rows):
                    if not headers:
                        raw_header = list(rows[header_idx])
                        headers = [str(c) if c is not None and str(c).strip() != "" else f"列{i+1}" for i, c in enumerate(raw_header)]
                        if cfg["include_source_sheet"]:
                            headers.insert(0, "来源工作表")
                    data_rows = rows[header_idx + 1:]
                else:
                    data_rows = rows

                total_scanned_rows += len(data_rows)

                for row in data_rows:
                    if not any(row):
                        continue

                    cells_to_check = []
                    row_len = len(row)
                    if target_cols is not None:
                        for idx in target_cols:
                            if idx < row_len and row[idx] is not None:
                                cells_to_check.append(str(row[idx]))
                    else:
                        for cell in row:
                            if cell is not None:
                                cells_to_check.append(str(cell))

                    if not cells_to_check:
                        if exclude_mode:
                            row_list = list(row)
                            if cfg["include_source_sheet"]:
                                row_list.insert(0, sheet_name)
                            extracted_items.append((-1, order_counter, row_list))
                            order_counter += 1
                        continue

                    is_matched = False
                    matched_kw_idx = -1

                    for kw_idx, kw in enumerate(keywords_processed):
                        hit = False
                        for cell_str in cells_to_check:
                            test_str = cell_str if case_sen else cell_str.lower()
                            if mode == "contains":
                                if kw in test_str:
                                    hit = True
                                    break
                            elif mode == "exact":
                                if kw == test_str:
                                    hit = True
                                    break
                            elif mode == "regex":
                                if regex_patterns[kw_idx].search(cell_str):
                                    hit = True
                                    break
                        if hit:
                            matched_kw_idx = kw_idx
                            break

                    if exclude_mode:
                        is_matched = (matched_kw_idx == -1)
                    else:
                        is_matched = (matched_kw_idx != -1)

                    if is_matched:
                        row_list = list(row)
                        if cfg["include_source_sheet"]:
                            row_list.insert(0, sheet_name)
                        extracted_items.append((matched_kw_idx, order_counter, row_list))
                        order_counter += 1

            if cfg["sort_by_kw"] and not exclude_mode:
                extracted_items.sort(key=lambda item: (item[0], item[1]))

            extracted_results = [item[2] for item in extracted_items]

            if cfg["remove_inner_empty"]:
                extracted_results = self.clean_inner_empty(extracted_results)
            
            if cfg["remove_empty_rows"]:
                extracted_results = [r for r in extracted_results if any(c is not None and str(c).strip() != "" for c in r)]

            if cfg["remove_duplicates"]:
                extracted_results = self.remove_duplicates(extracted_results)

            self.msg_queue.put(("SUCCESS", extracted_results, headers, total_scanned_rows))

        except Exception as e:
            error_details = traceback.format_exc()
            self.msg_queue.put(("ERROR", str(e), error_details))

    def read_all_sheets_data(self, file_path, sheets_to_read):
        data_dict = {}
        wb = None
        try:
            if file_path.endswith(".xlsx"):
                wb = load_workbook(file_path, read_only=True, data_only=True)
                for sheet_name in sheets_to_read:
                    if sheet_name in wb.sheetnames:
                        sheet = wb[sheet_name]
                        sheet_rows = []
                        consecutive_empty = 0
                        for row_tuple in sheet.iter_rows(values_only=True):
                            if not any(row_tuple):
                                consecutive_empty += 1
                                if consecutive_empty > 100:
                                    break
                            else:
                                consecutive_empty = 0
                            sheet_rows.append(row_tuple)
                        data_dict[sheet_name] = sheet_rows
            else:
                wb = xlrd.open_workbook(file_path)
                for sheet_name in sheets_to_read:
                    if sheet_name in wb.sheet_names():
                        sheet = wb.sheet_by_name(sheet_name)
                        data_dict[sheet_name] = [sheet.row_values(r) for r in range(sheet.nrows)]
        finally:
            if wb and hasattr(wb, 'close'):
                try: wb.close()
                except Exception: pass
        return data_dict

    def clean_inner_empty(self, data):
        cleaned = []
        for row in data:
            cleaned_row = []
            for cell in row:
                if isinstance(cell, str):
                    lines = [line.strip() for line in cell.splitlines() if line.strip()]
                    cleaned_row.append("\n".join(lines))
                else: cleaned_row.append(cell)
            cleaned.append(cleaned_row)
        return cleaned

    def remove_duplicates(self, data):
        seen = set()
        unique_data = []
        for row in data:
            row_tuple = tuple(str(c) if c is not None else "" for c in row)
            if row_tuple not in seen:
                seen.add(row_tuple)
                unique_data.append(row)
        return unique_data

    def process_queue(self):
        try:
            while True:
                msg = self.msg_queue.get_nowait()
                msg_type = msg[0]

                if msg_type == "SAFE_DROP_EVENT":
                    dropped_file = msg[1]
                    if self.is_processing:
                        messagebox.showwarning("警告", "后台正在提取数据，请勿切换文件！")
                    else:
                        ext = os.path.splitext(dropped_file)[1].lower()
                        if ext in ['.xlsx', '.xls']:
                            self.load_excel_file(dropped_file)
                        else:
                            messagebox.showwarning("格式不支持", "仅支持拖入 Excel 文件 (*.xlsx / *.xls)！")

                elif msg_type == "SUCCESS":
                    self.progress_bar.stop()
                    self.is_processing = False
                    self.btn_run.config(state=tk.NORMAL)

                    result = msg[1]
                    self.extracted_headers = msg[2]
                    total_scanned = msg[3]
                    self.extracted_data = result
                    count = len(result)
                    
                    ratio = (count / total_scanned * 100) if total_scanned > 0 else 0
                    self.stat_scanned_var.set(f"{total_scanned:,}")
                    self.stat_matched_var.set(f"{count:,}")
                    self.stat_ratio_var.set(f"{ratio:.1f}%")

                    if count > 0:
                        self.status_var.set("提取完成！已生成结果预览")
                        self.status_label.configure(bootstyle="success")
                        self.update_treeview_preview(result, self.extracted_headers)
                        self.save_btn.config(state=tk.NORMAL)
                    else:
                        self.status_var.set("扫描完毕：无匹配结果")
                        self.status_label.configure(bootstyle="warning")
                        self.save_btn.config(state=tk.DISABLED)

                elif msg_type == "ERROR":
                    self.progress_bar.stop()
                    self.is_processing = False
                    self.btn_run.config(state=tk.NORMAL)

                    err_msg, err_details = msg[1], msg[2]
                    self.status_var.set("提取失败，发生异常！")
                    self.status_label.configure(bootstyle="danger")
                    messagebox.showerror("提取失败", f"处理过程中发生严重异常:\n{err_msg}")
                    print(f"Extraction Error Details:\n{err_details}")

        except queue.Empty:
            pass
        finally:
            self.root.after(80, self.process_queue)

    def update_treeview_preview(self, data, headers):
        self.tree.delete(*self.tree.get_children())
        if not data: return

        max_cols_in_data = max(len(row) for row in data[:100])
        display_col_count = min(max_cols_in_data, 40)

        cols = [f"col_{i}" for i in range(display_col_count)]
        self.tree["columns"] = cols

        for i in range(display_col_count):
            if headers and i < len(headers):
                col_name = headers[i]
            else:
                col_name = f"列 {i+1}"
            self.tree.heading(f"col_{i}", text=col_name)
            self.tree.column(f"col_{i}", width=125, minwidth=90, anchor=tk.CENTER)

        for row in data[:100]:
            display_row = [str(c) if c is not None else "" for c in row[:display_col_count]]
            self.tree.insert("", tk.END, values=display_row)

    def generate_smart_filename(self, ext=".xlsx"):
        if self.file_path:
            src_dir = os.path.dirname(self.file_path)
            src_stem = os.path.splitext(os.path.basename(self.file_path))[0]
        else:
            src_dir = os.path.expanduser("~/Desktop")
            src_stem = "Excel数据"

        action_name = "过滤结果" if self.exclude_mode.get() else "提取结果"
        today_str = datetime.datetime.now().strftime("%Y%m%d")

        base_name = f"{src_stem}_{action_name}_{today_str}"
        candidate_name = f"{base_name}{ext}"

        counter = 1
        while os.path.exists(os.path.join(src_dir, candidate_name)):
            candidate_name = f"{base_name} ({counter}){ext}"
            counter += 1

        return src_dir, candidate_name

    def save_to_file(self):
        """
        🟢 纯净导出：清洗破坏 Excel XML 的非法不可见控制字符，保留原生类型
        """
        if not self.extracted_data: return

        initial_dir, default_filename = self.generate_smart_filename(".xlsx")

        save_path = filedialog.asksaveasfilename(
            initialdir=initial_dir,
            initialfile=default_filename,
            defaultextension=".xlsx",
            filetypes=[("Excel 工作簿 (*.xlsx)", "*.xlsx"), ("CSV 文件 (*.csv)", "*.csv")],
            title="导出提取结果"
        )
        if not save_path: return

        try:
            if save_path.lower().endswith(".csv"):
                with open(save_path, "w", newline="", encoding="utf-8-sig") as f:
                    writer = csv.writer(f)
                    if self.extracted_headers:
                        writer.writerow([sanitize_excel_value(h) for h in self.extracted_headers])
                    for row in self.extracted_data:
                        writer.writerow([sanitize_excel_value(c) for c in row])
            else:
                wb = Workbook()
                ws = wb.active
                ws.title = "提取结果"

                if self.extracted_headers:
                    ws.append([sanitize_excel_value(h) for h in self.extracted_headers])

                for row in self.extracted_data:
                    ws.append([sanitize_excel_value(cell) for cell in row])

                wb.save(save_path)
                wb.close()

            self.status_var.set(f"数据已保存: {os.path.basename(save_path)}")
            self.status_label.configure(bootstyle="success")

            if messagebox.askyesno("导出成功", f"文件保存成功！共导出 {len(self.extracted_data):,} 行数据。\n\n是否立即打开该文件？"):
                self.open_file_externally(save_path)

        except Exception as e:
            messagebox.showerror("保存失败", f"导出过程中发生错误:\n{str(e)}")

    def open_file_externally(self, filepath):
        try:
            if platform.system() == "Windows":
                os.startfile(filepath)
            elif platform.system() == "Darwin":
                subprocess.Popen(["open", filepath])
            else:
                subprocess.Popen(["xdg-open", filepath])
        except Exception as e:
            messagebox.showwarning("打开失败", f"无法自动打开文件:\n{str(e)}")

    def on_closing(self):
        if self.is_processing:
            if messagebox.askokcancel("退出确认", "程序正在后台提取数据，确定要退出吗？"):
                self.root.destroy()
        else:
            self.root.destroy()

if __name__ == "__main__":
    root = ttk.Window(themename="cosmo")
    app = ExcelExtractorApp(root)
    root.mainloop()