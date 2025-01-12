import tkinter as tk
from tkinter import filedialog
from openpyxl import load_workbook, Workbook
import xlrd
import ttkbootstrap as ttk
from ttkbootstrap.constants import *
from ttkbootstrap.tooltip import ToolTip

class ExcelExtractorApp:
    def __init__(self, root):
        self.root = root
        self.root.title("表格提取器 ExcelDataMatcher v 2.1  - QwejayHuang")
        self.root.geometry("640x520")

        # 初始化变量
        self.file_path = None
        self.sheet_names = []
        self.selected_sheet = tk.StringVar()
        self.search_all_sheets = tk.BooleanVar()
        self.header_row = tk.StringVar(value="1")
        self.no_header = tk.BooleanVar(value=True)
        self.remove_duplicates_var = tk.BooleanVar(value=False)
        self.remove_inner_empty_var = tk.BooleanVar(value=True)  # 单元格内空行
        self.remove_empty_rows_var = tk.BooleanVar(value=False)  # 整行空行

        self.create_widgets()
        self.create_context_menu()
        self.toggle_no_header()

    def create_widgets(self):
        main_frame = ttk.Frame(self.root, padding=10)
        main_frame.pack(fill=tk.BOTH, expand=True)

        # 文件选择部分
        file_frame = ttk.Frame(main_frame)
        file_frame.pack(fill=tk.X, pady=5)
        self.file_label = ttk.Label(file_frame, text="未选择文件", width=40, anchor=tk.W)
        self.file_label.pack(side=tk.LEFT, padx=5)
        ttk.Button(file_frame, text="选择表格文件", command=self.select_file).pack(side=tk.RIGHT, padx=5)

        # 工作表选择
        sheet_frame = ttk.Frame(main_frame)
        sheet_frame.pack(fill=tk.X, pady=5)
        ttk.Label(sheet_frame, text="选择工作表:", width=12, anchor=tk.E).pack(side=tk.LEFT, padx=5)
        self.sheet_combobox = ttk.Combobox(sheet_frame, textvariable=self.selected_sheet, state='readonly', width=20)
        self.sheet_combobox.pack(side=tk.LEFT, padx=5, fill=tk.X, expand=True)
        ttk.Checkbutton(sheet_frame, text="搜索所有", variable=self.search_all_sheets, command=self.toggle_search_all).pack(side=tk.RIGHT, padx=5)

        # 表头设置
        header_frame = ttk.Frame(main_frame)
        header_frame.pack(fill=tk.X, pady=5)
        ttk.Label(header_frame, text="表头行:", width=12, anchor=tk.E).pack(side=tk.LEFT, padx=5)
        self.header_entry = ttk.Entry(header_frame, textvariable=self.header_row, width=5)
        self.header_entry.pack(side=tk.LEFT, padx=5)
        ttk.Checkbutton(header_frame, text="不需要表头", variable=self.no_header, command=self.toggle_no_header).pack(side=tk.LEFT, padx=5)

        # 数据处理选项
        process_frame = ttk.LabelFrame(main_frame, text="数据处理选项", padding=10)
        process_frame.pack(fill=tk.X, pady=5)
        
        # 去除单元格内空行
        inner_empty_check = ttk.Checkbutton(
            process_frame,
            text="去除单元格内空行",
            variable=self.remove_inner_empty_var
        )
        inner_empty_check.pack(side=tk.LEFT, padx=10)
        ToolTip(inner_empty_check, text="清除单元格内容中的空行")  # 添加 ToolTip

        # 移除空行（整行）
        empty_rows_check = ttk.Checkbutton(
            process_frame,
            text="移除空行（整行）",
            variable=self.remove_empty_rows_var
        )
        empty_rows_check.pack(side=tk.LEFT, padx=10)
        ToolTip(empty_rows_check, text="删除所有单元格都为空的整行")  # 添加 ToolTip

        # 去除重复项
        duplicate_check = ttk.Checkbutton(
            process_frame,
            text="去除重复项",
            variable=self.remove_duplicates_var
        )
        duplicate_check.pack(side=tk.RIGHT, padx=10)
        ToolTip(duplicate_check, text="删除重复的数据行")  # 添加 ToolTip

        # 列输入区域
        ttk.Label(main_frame, text="输入列名（每列占一行）:", anchor=tk.W).pack(fill=tk.X, pady=5)
        self.column_text = tk.Text(main_frame, height=6, width=50)
        self.column_text.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)

        # 操作按钮
        ttk.Button(main_frame, text="提取并保存", command=self.extract_and_save, bootstyle=SUCCESS).pack(pady=10)

        # 状态栏
        self.status_var = tk.StringVar()
        self.status_var.set("就绪")
        status_bar = ttk.Label(main_frame, textvariable=self.status_var, anchor=tk.W)
        status_bar.pack(fill=tk.X, pady=5)

    def create_context_menu(self):
        self.context_menu = tk.Menu(self.root, tearoff=0)
        self.context_menu.add_command(label="粘贴", command=self.paste_text)
        self.context_menu.add_command(label="复制", command=self.copy_text)
        self.context_menu.add_command(label="清除", command=self.clear_text)
        self.column_text.bind("<Button-3>", self.show_context_menu)

    def show_context_menu(self, event):
        self.context_menu.post(event.x_root, event.y_root)

    def paste_text(self):
        self.column_text.event_generate("<<Paste>>")
        self.status_var.set("已粘贴内容")

    def copy_text(self):
        self.column_text.event_generate("<<Copy>>")
        self.status_var.set("已复制选中文本")

    def clear_text(self):
        self.column_text.delete("1.0", tk.END)
        self.status_var.set("已清除文本框内容")

    def select_file(self):
        file_path = filedialog.askopenfilename(filetypes=[("Excel files", "*.xlsx *.xls")])
        if file_path:
            self.file_path = file_path
            self.file_label.config(text=f"已选择文件: {file_path.split('/')[-1]}")
            try:
                if file_path.endswith(".xlsx"):
                    wb = load_workbook(file_path, read_only=True)
                    self.sheet_names = wb.sheetnames
                else:
                    wb = xlrd.open_workbook(file_path)
                    self.sheet_names = wb.sheet_names()
                self.sheet_combobox['values'] = self.sheet_names
                self.selected_sheet.set(self.sheet_names[0] if self.sheet_names else "")
                self.status_var.set("文件加载成功")
            except Exception as e:
                self.status_var.set(f"文件加载失败: {str(e)}")
                self.file_path = None
        else:
            self.status_var.set("操作取消: 未选择文件")

    def toggle_search_all(self):
        self.sheet_combobox.config(state='disabled' if self.search_all_sheets.get() else 'readonly')

    def toggle_no_header(self):
        if self.no_header.get():
            self.header_entry.config(state=tk.DISABLED)
            self.header_row.set("None")
        else:
            self.header_entry.config(state=tk.NORMAL)
            self.header_row.set("1")

    def extract_and_save(self):
        if not self.file_path:
            self.status_var.set("错误: 请先选择表格文件")
            return

        column_names = self.column_text.get("1.0", "end-1c").splitlines()
        column_names = [col.strip() for col in column_names if col.strip()]

        if not column_names:
            self.status_var.set("错误: 列名不能为空")
            return

        header_row = self.header_row.get()
        if header_row == "None" or self.no_header.get():
            header_row = None
        else:
            try:
                header_row = int(header_row) - 1
            except ValueError:
                self.status_var.set("错误: 表头行必须是整数或选择不需要表头")
                return

        try:
            all_data = self.get_all_data(header_row)
            
            # 分步处理数据
            if self.remove_inner_empty_var.get():
                all_data = self.remove_inner_empty_lines(all_data)
                self.status_var.set("已处理单元格内空行")

            if self.remove_empty_rows_var.get():
                all_data = self.remove_empty_rows(all_data)
                self.status_var.set("已移除空行（整行）")

            extracted_data = self.find_matching_rows(all_data, column_names)
            
            if self.remove_duplicates_var.get():
                extracted_data = self.remove_duplicates(extracted_data)
                self.status_var.set("数据已去重处理")

            if not extracted_data:
                self.status_var.set("信息: 未找到任何匹配的行")
                return

            save_path = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel files", "*.xlsx *.xls")])
            if save_path:
                wb = Workbook()
                ws = wb.active
                for row in extracted_data:
                    ws.append(row)
                wb.save(save_path)
                self.status_var.set(f"成功: 数据已保存到 {save_path}")
            else:
                self.status_var.set("操作取消: 未选择保存路径")
        except Exception as e:
            self.status_var.set(f"错误: 提取或保存失败: {e}")

    def get_all_data(self, header_row):
        all_data = []
        if self.file_path.endswith(".xlsx"):
            workbook = load_workbook(self.file_path, read_only=True)
            if self.search_all_sheets.get():
                for sheet_name in self.sheet_names:
                    sheet = workbook[sheet_name]
                    data = list(sheet.iter_rows(values_only=True))
                    if header_row is not None:
                        data = data[header_row + 1:]
                    all_data.extend(data)
            else:
                sheet = workbook[self.selected_sheet.get()]
                data = list(sheet.iter_rows(values_only=True))
                if header_row is not None:
                    data = data[header_row + 1:]
                all_data.extend(data)
        elif self.file_path.endswith(".xls"):
            workbook = xlrd.open_workbook(self.file_path)
            if self.search_all_sheets.get():
                for sheet_name in self.sheet_names:
                    sheet = workbook.sheet_by_name(sheet_name)
                    data = [sheet.row_values(row) for row in range(sheet.nrows)]
                    if header_row is not None:
                        data = data[header_row + 1:]
                    all_data.extend(data)
            else:
                sheet = workbook[self.selected_sheet.get()]
                data = [sheet.row_values(row) for row in range(sheet.nrows)]
                if header_row is not None:
                    data = data[header_row + 1:]
                all_data.extend(data)
        return all_data

    def find_matching_rows(self, all_data, column_names):
        extracted_data = []
        column_names_lower = [col.strip().lower() for col in column_names]
        for row in all_data:
            for cell in row:
                cell_str = str(cell).lower()
                if any(col in cell_str for col in column_names_lower):
                    extracted_data.append(row)
                    break
        return extracted_data

    def remove_duplicates(self, data):
        seen = set()
        unique_data = []
        for row in data:
            row_tuple = tuple(row)
            if row_tuple not in seen:
                seen.add(row_tuple)
                unique_data.append(row)
        return unique_data

    def remove_inner_empty_lines(self, data):
        """处理单元格内部空行"""
        cleaned_data = []
        for row in data:
            cleaned_row = []
            for cell in row:
                if isinstance(cell, str):
                    lines = cell.splitlines()
                    non_empty_lines = [line for line in lines if line.strip()]
                    cleaned_cell = '\n'.join(non_empty_lines)
                    cleaned_row.append(cleaned_cell)
                else:
                    cleaned_row.append(cell)
            cleaned_data.append(cleaned_row)
        return cleaned_data

    def remove_empty_rows(self, data):
        """移除整行空数据"""
        return [row for row in data if any(
            str(cell).strip() if isinstance(cell, str) else cell is not None 
            for cell in row
        )]

if __name__ == "__main__":
    root = ttk.Window(themename="superhero")
    app = ExcelExtractorApp(root)
    root.mainloop()